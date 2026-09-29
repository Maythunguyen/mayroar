"""Shared test helpers."""

from __future__ import annotations

import shutil
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook

FIXTURES = Path(__file__).parent / "fixtures"

AFCD_PROFILE_HEADERS = [
    "Public Food Key",
    "Classification",
    "Derivation",
    "Food Name",
    "Energy with dietary fibre, equated \n(kJ)",
    "Protein \n(g)",
    "Fat, total \n(g)",
    "Available carbohydrate, with sugar alcohols \n(g)",
    "Total dietary fibre \n(g)",
    "Total sugars (g)",
    "Alcohol \n(g)",
    "Sodium (Na) \n(mg)",
    "Calcium (Ca) \n(mg)",
    "Iron (Fe) \n(mg)",
    "Magnesium (Mg) \n(mg)",
    "Vitamin D3 equivalents \n(ug)",
]
# key, class, derivation, name, kJ, protein, fat, carbs, fibre, sugars, alcohol, Na, Ca, Fe, Mg, vit D
AFCD_ROWS = [
    [
        "F000001",
        1,
        "Analysed",
        "Chicken, breast, lean flesh, raw",
        412,
        22.5,
        0.8,
        0,
        0,
        0,
        0,
        51,
        5,
        0.4,
        28,
        0.1,
    ],
    [
        "F000002",
        2,
        "Analysed",
        "Milk, cow, fluid, regular fat",
        272,
        3.4,
        3.5,
        4.7,
        0,
        4.7,
        0,
        38,
        115,
        0,
        11,
        None,
    ],
    ["F000003", 3, "Label data", "Butter, salt reduced", 3000, 0.6, 81, 0.6, 0, 0.6, 0, 300, 15, 0, 2, 0.5],
]


def _sheet_with_title(wb: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(["Australian Food Composition Database - Release 3"])  # FSANZ puts a title first
    ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)


def build_afcd_files(folder: Path) -> None:
    """Writes a tiny copy of the two AFCD files, shaped like the real ones."""
    folder.mkdir(parents=True, exist_ok=True)
    profiles = Workbook()
    profiles.remove(profiles.active)  # type: ignore[arg-type]
    _sheet_with_title(profiles, "All solids & liquids per 100 g", AFCD_PROFILE_HEADERS, AFCD_ROWS)
    _sheet_with_title(profiles, "Liquids only per 100 mL", AFCD_PROFILE_HEADERS, [AFCD_ROWS[1]])
    profiles.save(folder / "AFCD_Release_3_-_Nutrient_profiles.xlsx")

    details = Workbook()
    details.remove(details.active)  # type: ignore[arg-type]
    _sheet_with_title(
        details,
        "Food details",
        ["Public Food Key", "Food Name", "Food Description", "Specific Gravity"],
        [
            ["F000001", "Chicken", "Breast, no skin", None],
            ["F000002", "Milk", "Regular fat cow's milk", 1.03],
            ["F000003", "Butter", "Salt reduced", None],
        ],
    )
    details.save(folder / "AFCD_Release_3_-_Food_Details.xlsx")


@pytest.fixture
def raw_dir(tmp_path: Path) -> Path:
    """A data/raw folder holding the small test versions of every source."""
    raw = tmp_path / "raw"
    shutil.copytree(FIXTURES / "usda_foundation", raw / "usda_foundation")
    build_afcd_files(raw / "fsanz_afcd")
    return raw


def _nutrient(name: str, per_100g: float | None) -> str:
    value = "null" if per_100g is None else repr(per_100g)
    return f"{{'name': '{name}', 'value': {value}, '100g': {value}, 'serving': null, 'unit': 'g'}}"


# code, names, brand, quantity, serving size, serving grams, per, nutrients (name, per 100 g)
OffProduct = tuple[str, str, str | None, str | None, str | None, str | None, str | None, list]
OFF_PRODUCTS: list[OffProduct] = [
    (
        "9300000000011",
        "[{'lang': 'main', 'text': 'Yoghurt'}, {'lang': 'en', 'text': 'Greek Yoghurt Plain'}]",
        "Brand A",
        "500 g",
        "170 g",
        "170",
        "100g",
        [
            ("proteins", 9.0),
            ("fat", 5.0),
            ("carbohydrates", 4.0),
            ("energy-kcal", 97.0),
            ("sodium", 0.05),
            ("calcium", 0.12),
            ("vitamin-d", 0.0000025),
            ("salt", 0.125),
        ],
    ),
    (
        "9300000000028",
        "[{'lang': 'main', 'text': 'Protein Bar'}]",
        "Brand B",
        "60 g",
        "1 bar (60 g)",
        "60",
        "serving",
        [("proteins", 30.0), ("fat", 10.0), ("carbohydrates", 30.0), ("energy-kj", 1500.0)],
    ),
    (
        "9300000000035",
        "[{'lang': 'main', 'text': 'Mystery Snack'}]",
        None,
        None,
        None,
        None,
        None,
        [("proteins", None)],
    ),
    (
        "abc123",
        "[{'lang': 'main', 'text': 'Bad Barcode'}]",
        None,
        None,
        None,
        None,
        "100g",
        [("proteins", 1.0)],
    ),
    (
        "9300000000042",
        "[{'lang': 'main', 'text': 'Cordial'}]",
        "Brand C",
        "1 L",
        None,
        "abc",
        None,
        [("proteins", 0.0), ("fat", 0.0), ("carbohydrates", 20.0), ("energy", 340.0)],
    ),
]


def build_off_extract(folder: Path, day: str = "2026-09-23") -> Path:
    """Writes a tiny extract shaped like the real open_food_facts_au file."""
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"open_food_facts_au_{day}.parquet"

    def text(value: str | None) -> str:
        return "null::varchar" if value is None else "'" + value.replace("'", "''") + "'"

    selects = []
    for code, names, brand, qty, size, grams, per, nutrients in OFF_PRODUCTS:
        nutriments = "[" + ", ".join(_nutrient(n, v) for n, v in nutrients) + "]"
        selects.append(
            f"""select '{code}' as code, {names} as product_name, {text(brand)} as brands,
                       {text(qty)} as quantity, {text(size)} as serving_size,
                       {text(grams)} as serving_quantity,
                       {text(per)} as nutrition_data_per, {nutriments} as nutriments,
                       ['en:australia'] as countries_tags, 0.8::float as completeness,
                       []::varchar[] as data_quality_errors_tags, 1700000000::bigint as last_modified_t,
                       false as obsolete"""
        )
    with duckdb.connect() as con:
        con.execute(f"copy ({' union all '.join(selects)}) to '{path.as_posix()}' (format parquet)")
    return path


@pytest.fixture
def make_off_extract():  # type: ignore[no-untyped-def]
    """Gives tests the build_off_extract helper."""
    return build_off_extract
